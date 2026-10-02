"""Builds writeup.pdf (2 pages) and slides.pptx from outputs/results.json + figures.
Run AFTER train.py:  python make_docs.py"""
import json, os
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle
from pptx import Presentation
from pptx.util import Inches, Pt

OUT = "outputs"
R = json.load(open(f"{OUT}/results.json"))
M, D = R["models"], R["dataset"]
names = list(M)
scratch, frozen, ft = (M[n] for n in names)
pct = lambda v: f"{100*v:.2f}%"
err = lambda m: m["cm"][0][1] + m["cm"][1][0]
gain = 100 * (ft["test_acc"] - scratch["test_acc"])
TEAM = "Team 7 | PES1UG24CS376 & PES1UG24CS409 | Section G | UE24CS352A Machine Learning"
TITLE = "Household Animals Classification Using Deep Learning"
fig = lambda f: f"{OUT}/{f}" if os.path.exists(f"{OUT}/{f}") else None

# ---------------- text shared by PDF and slides ----------------
problem = ("Automatically identify household animals (cat vs dog) from photographs using deep learning. "
           "Accurate species classification is the first step towards monitoring pets' behaviour and health. "
           "We follow the CS229 (Stanford, 2020) reference by Lei Lin and reproduce its pipeline in PyTorch.")
dataset = (f"Dogs vs Cats (Kaggle), obtained through the Hugging Face mirror microsoft/cats_vs_dogs (corrupted files already removed). "
           f"{D['total']} labelled RGB images, two balanced classes, variable size and quality. Random split (seed 42): "
           f"{D['train']} train / {D['val']} validation / {D['test']} test. Images resized (128x128 for the scratch CNN, 224x224 for VGG-16) "
           "and normalised with ImageNet statistics. Training augmentation: horizontal flip, rotation (15 deg), colour jitter.")
approach = [
    f"Baseline: VGG-style CNN trained from scratch (4 blocks of 3x3 conv + ReLU + max-pool with 32/64/128/128 filters, dropout 0.5, Adam lr 1e-3, {scratch['epochs']} epochs).",
    f"Transfer learning: ImageNet-pretrained VGG-16, conv layers frozen, new classifier (25088-256-1) trained for {frozen['epochs']} epochs (lr 1e-3).",
    f"Fine-tuning: block5 (conv5_1 onward) unfrozen and trained with lr 1e-5 for {ft['epochs']} epochs.",
    "Interpretability: intermediate feature-map visualisation (scratch CNN) and Grad-CAM heatmaps (VGG-16).",
]
impl = ("Single script train.py (PyTorch, Google Colab GPU). Binary classification with one logit and BCE-with-logits loss, mixed-precision training, "
        "best checkpoint chosen on validation accuracy, final numbers reported on the held-out test set (accuracy, F1, confusion matrix). "
        "make_docs.py generates this report and the slides from the saved results.")
concl = [
    f"Best model: {names[2]} with {pct(ft['test_acc'])} test accuracy (F1 {ft['test_f1']:.3f}), versus {pct(scratch['test_acc'])} for the scratch CNN "
    f"({'+' if gain>=0 else ''}{gain:.2f} points).",
    f"Fine-tuning improved only slightly over the frozen VGG-16 ({pct(frozen['test_acc'])} to {pct(ft['test_acc'])}; {err(frozen)} vs {err(ft)} errors on {D['test']} test images). "
    "This small gap is within what a single run could produce, so most of the gain comes from the pretrained features themselves.",
    "Pretrained ImageNet features transfer well to cat/dog classification; fine-tuning the last conv block adjusts them to the task at a low learning rate.",
    "Grad-CAM shows where the network looks (e.g. face and body of the animal), which helps check that predictions rely on the animal rather than the background.",
    "Limitations: only two classes and one dataset; no real-world (non-curated) images. Future work: more species/breeds, other backbones (ResNet, MobileNet), pet-health detection.",
]
refs = ["L. Lin, Household Animals Classification Using Deep Learning, CS229, Stanford, 2020.",
        "K. Simonyan, A. Zisserman, Very Deep Convolutional Networks for Large-Scale Image Recognition, 2014.",
        "R. Selvaraju et al., Grad-CAM: Visual Explanations from Deep Networks, ICCV 2017.",
        "Kaggle Dogs vs Cats: kaggle.com/c/dogs-vs-cats"]
rows = [["Model", "Trainable params", "Best val acc", "Test acc", "Test F1"]] + [
    [n, f"{M[n]['trainable_params']:,}", pct(M[n]["best_val_acc"]), pct(M[n]["test_acc"]), f"{M[n]['test_f1']:.3f}"] for n in names]

# ---------------- PDF ----------------
ss = getSampleStyleSheet()
B = ParagraphStyle("b", parent=ss["Normal"], fontSize=8.8, leading=11, spaceAfter=2)
H = ParagraphStyle("h", parent=ss["Heading3"], fontSize=10.5, spaceBefore=5, spaceAfter=2)
T_ = ParagraphStyle("t", parent=ss["Title"], fontSize=15, spaceAfter=2)
doc = SimpleDocTemplate("writeup.pdf", pagesize=A4, leftMargin=1.6*cm, rightMargin=1.6*cm, topMargin=1.3*cm, bottomMargin=1.2*cm)
s = [Paragraph(TITLE, T_), Paragraph(TEAM, ParagraphStyle("c", parent=B, alignment=1)), Spacer(1, 4)]
s += [Paragraph("1. Problem Statement", H), Paragraph(problem, B), Paragraph("2. Dataset", H), Paragraph(dataset, B), Paragraph("3. Approach", H)]
s += [Paragraph("&bull; " + a, B) for a in approach]
s += [Paragraph("4. Implementation Overview", H), Paragraph(impl, B), Paragraph("5. Results", H)]
t = Table(rows, colWidths=[5.3*cm, 3.3*cm, 3*cm, 2.8*cm, 2.5*cm])
t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8), ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("GRID", (0, 0), (-1, -1), .4, colors.grey)]))
s += [t, Spacer(1, 4)]
if fig("curves.png"): s.append(Image(fig("curves.png"), width=15*cm, height=15*cm*3.8/11))
if fig("gradcam.png"): s.append(Image(fig("gradcam.png"), width=15.5*cm, height=15.5*cm*4.3/12))
s += [Paragraph("6. Conclusions", H)] + [Paragraph("&bull; " + c, B) for c in concl]
s += [Paragraph("References", H)] + [Paragraph(f"[{i+1}] {r}", ParagraphStyle("r", parent=B, fontSize=7.5, leading=9)) for i, r in enumerate(refs)]
doc.build(s)

# ---------------- Slides ----------------
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)

def slide(title, bullets=None, img=None, size=20):
    sl = prs.slides.add_slide(prs.slide_layouts[5]); sl.shapes.title.text = title
    sl.shapes.title.text_frame.paragraphs[0].font.size = Pt(34)
    if bullets:
        w = Inches(6.2) if img else Inches(12.3)
        tf = sl.shapes.add_textbox(Inches(0.5), Inches(1.5), w, Inches(5.5)).text_frame; tf.word_wrap = True
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = "• " + b; p.font.size = Pt(size); p.space_after = Pt(10)
    if img: sl.shapes.add_picture(img, Inches(6.9) if bullets else Inches(0.8), Inches(1.6), width=Inches(6.0) if bullets else Inches(11.7))
    return sl

sl = prs.slides.add_slide(prs.slide_layouts[0]); sl.shapes.title.text = TITLE
sl.placeholders[1].text = "UE24CS352A Machine Learning - Mini Project\nTeam 7 | PES1UG24CS376, PES1UG24CS409 | Section G"
slide("Problem Statement", [problem])
slide("Dataset", [dataset], size=18)
slide("Approach", approach, size=18)
slide("Implementation", [impl, "Stack: PyTorch, torchvision, Hugging Face datasets, scikit-learn, Matplotlib; Google Colab GPU.", "Demo: run train.py live (or load saved checkpoint) and show outputs."], size=18)
sl = slide("Results", None)
tb = sl.shapes.add_table(len(rows), 5, Inches(0.5), Inches(1.6), Inches(12.3), Inches(2.2)).table
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        tb.cell(i, j).text = v; tb.cell(i, j).text_frame.paragraphs[0].font.size = Pt(16)
if fig("confusion_matrix.png"): sl.shapes.add_picture(fig("confusion_matrix.png"), Inches(0.8), Inches(4.0), height=Inches(3.2))
if fig("curves.png"): sl.shapes.add_picture(fig("curves.png"), Inches(4.4), Inches(4.0), height=Inches(3.2))
if fig("curves.png"): pass
if fig("activations.png"): slide("What the CNN learns: feature maps", img=fig("activations.png"))
if fig("gradcam.png"): slide("Grad-CAM heatmaps (VGG-16)", img=fig("gradcam.png"))
slide("Conclusions & Future Work", concl, size=18)
slide("Thank You - Questions?", [TEAM])
prs.save("slides.pptx"); print("Created writeup.pdf and slides.pptx")
